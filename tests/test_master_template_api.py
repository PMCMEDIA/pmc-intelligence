import os
import tempfile
import unittest
from io import BytesIO
from pptx import Presentation
from unittest.mock import patch
_tmp=tempfile.TemporaryDirectory()
os.environ.setdefault('DATABASE_PATH',os.path.join(_tmp.name,'template-api.db'))
os.environ.setdefault('SECRET_KEY','test-only-master-template')
from app import app
from storage import get
from workspace_core import revision

class TemplateApiTests(unittest.TestCase):
    def setUp(self):
        app.config['TESTING']=True
        self.client=app.test_client()
        with self.client.session_transaction() as sess:sess['owner']='test@pmcne.com'
        with patch('app.research_url',return_value={'ok':False,'error':'Offline fixture'}):
            response=self.client.post('/api/projects',json={'client':'Harbor API Fixture','goal':'More qualified inquiries'})
        self.assertEqual(response.status_code,200)
        self.pid=response.json['id'];self.path='/api/projects/'+self.pid+'/workspace'
    def saved(self):return get('test@pmcne.com',self.pid)
    def test_master_status_requires_login(self):
        self.assertEqual(app.test_client().get('/api/strategy-template').status_code,401)
        self.assertTrue(self.client.get('/api/strategy-template').json['ready'])
    def test_draft_edit_preview_export_and_restore(self):
        response=self.client.post(self.path+'/deck',json={'revision':revision(self.saved())})
        self.assertEqual(response.status_code,200,response.json)
        original=response.json['project']['deck_outline']
        self.assertEqual(original[0]['template_key'],'cover')
        original[0]['fields']['title']='Verified manual edit\nMarketing Proposal'
        response=self.client.post(self.path,json={'revision':revision(self.saved()),'patch':{'deck_outline':original}})
        self.assertEqual(response.status_code,200,response.json)
        self.assertIn('Verified manual edit',response.json['project']['deck_outline'][0]['body'])
        download=self.client.get('/api/projects/'+self.pid+'/deck.pptx')
        self.assertEqual(download.status_code,200)
        prs=Presentation(BytesIO(download.data))
        self.assertIn('Verified manual edit',' '.join(x.text for x in prs.slides[0].shapes if x.has_text_frame))
        response=self.client.post(self.path+'/deck',json={'revision':revision(self.saved()),'rebuild':True})
        self.assertEqual(response.status_code,200)
        self.assertIn('Verified manual edit',self.saved()['previous_deck_outline'][0]['body'])
        response=self.client.post(self.path+'/restore-deck',json={'revision':revision(self.saved())})
        self.assertEqual(response.status_code,200)
        self.assertIn('Verified manual edit',self.saved()['deck_outline'][0]['body'])
    def test_bad_fields_do_not_change_saved_project(self):
        response=self.client.post(self.path+'/deck',json={'revision':revision(self.saved())})
        slides=response.json['project']['deck_outline'];slides[0]['fields']['title']='X'*181
        before=self.saved()
        response=self.client.post(self.path,json={'revision':revision(before),'patch':{'deck_outline':slides}})
        self.assertEqual(response.status_code,400)
        self.assertEqual(self.saved(),before)

if __name__=='__main__':unittest.main(verbosity=2)
