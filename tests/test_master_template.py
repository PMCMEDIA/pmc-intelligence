"""Real-master regressions: no network, no live projects, no LLM claims."""
from copy import deepcopy
from io import BytesIO
import posixpath
import unittest
from zipfile import ZipFile
from lxml import etree
from pptx import Presentation
from pmc_template_library import (LAYOUTS, NS, VERSION, make_slide, preview_text,
                                  render_master, source_parts, template_status, validate_slide)
from proposal_blueprint import default_outline, required_assets, channel_names
from workspace_core import apply_patch
from strategy_engine import build_strategy


def fixture():
    p={'client':'Harbor Sample Company','goal':'Increase qualified service inquiries',
       'industry':'Other','budget':'$12,000 monthly','workflow_version':2}
    p['strategy']=build_strategy(p,{'ok':False,'error':'Test fixture'})
    p['website_pricing']={'pages':12,'page_rate':850,'discovery':500,'templates':3,
                          'mockup_rate':400,'technical':700,'adjustment':200}
    p['investment_allocations']=[{'department':'Digital Media','amount':5000,'percent':100}]
    return p

class MasterTemplateTests(unittest.TestCase):
    def test_pinned_source_and_library(self):
        self.assertTrue(template_status()['ready'])
        self.assertEqual(template_status()['reference_slides'],86)
        self.assertGreaterEqual(len(LAYOUTS),18)

    def test_template_outline_uses_source_order_and_all_departments(self):
        slides=default_outline(fixture());keys=[s['template_key'] for s in slides]
        self.assertEqual(keys[:4],['cover','positioning','agency_locations','agency_approach'])
        self.assertEqual(keys[-1],'closing')
        self.assertIn('seo',keys);self.assertIn('seo_execution',keys)
        self.assertLess(keys.index('search'),keys.index('seo'))
        self.assertLess(keys.index('seo'),keys.index('website'))
        self.assertTrue(all(s['template_version']==VERSION for s in slides))

    def test_only_selected_departments_and_channels(self):
        p=fixture();p['strategy']['departments']={'Digital Media':['Google Search']}
        p.pop('website_pricing')
        p['tactic_allocations']={'Digital Media':[{'name':'Google Search','percent':100},{'name':'Performance Max','percent':0}]}
        slides=default_outline(p);keys=[s['template_key'] for s in slides]
        self.assertIn('search',keys)
        for key in ['pmax','paid_social','ctv','seo','website','creative']:self.assertNotIn(key,keys)

    def test_production_assets_are_channel_driven_not_priced(self):
        p=fixture();assets='\n'.join(required_assets(p))
        self.assertIn(':30 commercial master and :15 cutdown',assets)
        self.assertIn(':15 vertical',assets);self.assertIn('do not count shared assets',assets)
        self.assertNotIn('$',assets)
        p['tactic_allocations']={'Digital Media':[{'name':'Google Search','percent':100}]}
        self.assertNotIn('CTV','\n'.join(required_assets(p)))

    def test_manual_strategy_values_flow_into_deck(self):
        p=fixture();p['department_details']={'SEO / AEO':{'objective':'Approved exact objective','kpis':'Approved exact measures'}}
        deck='\n'.join(s['body'] for s in default_outline(p))
        self.assertIn('Approved exact objective',deck);self.assertIn('Approved exact measures',deck)

    def test_website_math_no_sample_prices_or_recurring_guess(self):
        deck=default_outline(fixture());web=next(s for s in deck if s['title']=='Website Investment')
        self.assertEqual(web['rows'][-1][2],'$12,400')
        marketing=next(s for s in deck if s['title']=='Marketing Investment')
        self.assertEqual(marketing['rows'][1][1],'Period to confirm')
        self.assertEqual(marketing['rows'][1][2],'$5,000')

    def test_master_export_opens_and_has_editable_elements(self):
        deck=default_outline(fixture());stream=render_master(deck);prs=Presentation(stream)
        self.assertEqual(len(prs.slides),len(deck));self.assertEqual(prs.slide_width,18288000)
        self.assertGreater(sum(sh.has_text_frame for slide in prs.slides for sh in slide.shapes),40)
        self.assertTrue(any(sh.has_table for slide in prs.slides for sh in slide.shapes))

    def test_no_source_clients_notes_fonts_thumbnails_or_orphan_parts(self):
        with ZipFile(render_master(default_outline(fixture()))) as package:
            names=set(package.namelist())
            for name in names:
                self.assertFalse(any(s in name for s in ('notesSlide','notesMaster','fonts/','comments/','changesInfo','thumbnail','embeddings/')),name)
                if name.endswith('.xml'):
                    visible=' '.join(etree.fromstring(package.read(name)).xpath('//*[local-name()="t"]/text()'))
                    for banned in ('DeQuattro','VIVA','Viva','Cranston','Botox','Morpheus8','Crow’s Nest','Il Massimo'):
                        self.assertNotIn(banned,visible,name)
                if name.endswith('.rels'):
                    parent='' if name=='_rels/.rels' else name.rsplit('/_rels/',1)[0]
                    for rel in etree.fromstring(package.read(name)):
                        self.assertNotEqual(rel.get('TargetMode'),'External')
                        target=posixpath.normpath(posixpath.join(parent,rel.get('Target'))).lstrip('/')
                        self.assertIn(target,names,(name,target))

    def test_no_screenshot_export_and_source_fixed_copy_retained(self):
        p=make_slide('positioning',{},'Agency')
        self.assertIn('This is where brand gravity is born.',p['body'])
        prs=Presentation(render_master([p]));self.assertTrue(any(sh.has_text_frame for sh in prs.slides[0].shapes))
        self.assertEqual(preview_text(p),p['body'])

    def test_duplicate_layouts_and_reordering(self):
        one=make_slide('executive',{'title':'FIRST','copy':'First copy.'})
        two=make_slide('executive',{'title':'SECOND','copy':'Second copy.'})
        hidden=make_slide('closing',{});hidden['include']=False
        prs=Presentation(render_master([two,one,hidden]))
        self.assertEqual(len(prs.slides),2)
        self.assertIn('SECOND',' '.join(s.text for s in prs.slides[0].shapes if s.has_text_frame))
        self.assertIn('First copy.',' '.join(s.text for s in prs.slides[1].shapes if s.has_text_frame))

    def test_slide_edit_preview_and_export_are_same_values(self):
        slide=make_slide('executive',{'title':'EXECUTIVE SUMMARY','copy':'Old value.'})
        slide['fields']['copy']='Edited value & <literal text> remains editable.'
        result=apply_patch({'deck_outline':[slide]}, {'deck_outline':[slide]})
        self.assertIn('Edited value & <literal text>',result['deck_outline'][0]['body'])
        prs=Presentation(render_master(result['deck_outline']))
        rendered=' '.join(sh.text for sh in prs.slides[0].shapes if sh.has_text_frame)
        self.assertIn('Edited value & <literal text>',rendered);self.assertNotIn('Old value.',rendered)

    def test_invalid_or_oversized_fields_and_mixed_decks_rejected(self):
        good=make_slide('executive',{'title':'EXECUTIVE SUMMARY','copy':'Valid copy.'})
        bad=deepcopy(good);bad['fields']['copy']='X'*1501
        with self.assertRaises(ValueError):validate_slide(bad)
        with self.assertRaises(ValueError):render_master([{'template_key':'../../etc/passwd'}])
        with self.assertRaises(ValueError):apply_patch({}, {'deck_outline':[good,{'title':'Legacy','body':'Do not silently replace'}]})

    def test_placeholder_forecasts_cannot_look_like_sourced_metrics(self):
        p=fixture();p['campaign_outlook']={'scenarios':[{'projected_revenue':999999,'roas':111.0}],'roi_ready':True}
        body='\n'.join(s['body'] for s in default_outline(p))
        self.assertNotIn('999999',body);self.assertNotIn('111.0',body)
        self.assertIn('Independently sourced industry benchmarks',body)

    def test_overfilled_table_cell_is_rejected_instead_of_clipped(self):
        with self.assertRaises(ValueError):
            make_slide('investment',{'title':'Investment'},rows=[['Scope','Period','Cost','Review']]+[['Valid','One-time','$500','X'*400] for _ in range(8)])

    def test_existing_deck_not_modified_by_building_a_plan(self):
        p=fixture();p['deck_outline']=[{'title':'Manual original','body':'Original slide edit.'}]
        before=deepcopy(p);default_outline(p)
        self.assertEqual(p,before)

    def test_long_department_copy_is_split_not_dropped(self):
        p=fixture();p['strategy']['departments']={'SEO / AEO':['Technical audit']}
        p['department_details']={'SEO / AEO':{'deliverables':('Preserve the exact requested detail. '*120)}}
        slides=default_outline(p);texts=' '.join(s['fields'].get('copy','') for s in slides)
        self.assertEqual(texts.count('Preserve the exact requested detail.'),120)

if __name__=='__main__':unittest.main(verbosity=2)
