from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.chart.data import ChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from io import BytesIO

NAVY=RGBColor(11,23,36); TEAL=RGBColor(22,183,200); MUTED=RGBColor(93,111,124); WHITE=RGBColor(255,255,255)

def default_outline(p):
    slides=[{"title":str(p.get("client","Client"))+" Strategy","body":"Integrated Marketing Strategy\nPrepared by PMC Media Group","type":"cover","include":True},{"title":"Executive Direction","body":p.get("strategy",{}).get("executive",""),"type":"content","include":True}]
    facts=p.get("strategy",{}).get("facts",[])[:6]
    slides.append({"title":"Research + Discovery","body":"\n".join(["• "+str(x.get("label"))+": "+str(x.get("value")) for x in facts]),"type":"content","include":True})
    for name,recs in p.get("strategy",{}).get("departments",{}).items():
        d=p.get("department_details",{}).get(name,{})
        body=[]
        if d.get("objective"): body.append("OBJECTIVE\n"+d["objective"])
        if recs: body.append("RECOMMENDATIONS\n"+"\n".join(["• "+x for x in recs]))
        if d.get("deliverables"): body.append("DELIVERABLES\n"+d["deliverables"])
        if d.get("kpis"): body.append("SUCCESS CRITERIA\n"+d["kpis"])
        slides.append({"title":name,"body":"\n\n".join(body),"type":"department","include":True})
    inv=p.get("investment_allocations",[])
    if inv: slides.append({"title":"Investment Framework","body":"\n".join(["• "+str(x.get("department"))+": "+str(x.get("percent",0))+"% | $"+format(float(x.get("amount",0)),",.0f") for x in inv]),"type":"content","include":True})
    tactics=p.get("tactic_allocations",{})
    if tactics:
        body=[]
        for dept,items in tactics.items():
            body.append(dept.upper())
            body.extend(["• "+str(x.get("name"))+" — "+str(x.get("percent",0))+"% of department | KPI: "+str(x.get("kpi") or "TBD") for x in items])
        slides.append({"title":"Channel + Tactic Plan","body":"\n".join(body),"type":"channels","include":True})
        slides.append({"title":"Measurement + Success Framework","body":"\n".join(["• "+str(x.get("name"))+": "+str(x.get("kpi") or "TBD") for items in tactics.values() for x in items]),"type":"measurement","include":True})
    slides.append({"title":"Next Steps","body":"• Confirm launch priorities and timing\n• Finalize production and implementation requirements\n• Activate measurement and reporting plan","type":"closing","include":True})
    return slides

def _rgb(hexv, fallback):
    try:
        h=(hexv or "").replace("#",""); return RGBColor(int(h[0:2],16),int(h[2:4],16),int(h[4:6],16))
    except: return fallback

def build_pptx(p):
    prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
    brand=p.get("deck_branding",{}); accent=_rgb(brand.get("accent_color"),TEAL)
    for item in (p.get("deck_outline") or default_outline(p)):
        if item.get("include",True) is False: continue
        slide=prs.slides.add_slide(prs.slide_layouts[6]); bg=slide.background.fill; bg.solid(); bg.fore_color.rgb=NAVY if item.get("type") in ("cover","closing") else WHITE
        accent=slide.shapes.add_shape(1,Inches(.65),Inches(.65),Inches(.12),Inches(.72)); accent.fill.solid(); accent.fill.fore_color.rgb=accent; accent.line.fill.background()
        title=slide.shapes.add_textbox(Inches(1),Inches(.62),Inches(11.5),Inches(.85)); p1=title.text_frame.paragraphs[0]; p1.text=item.get("title",""); p1.font.size=Pt(34 if item.get("type")=="cover" else 28); p1.font.bold=True; p1.font.color.rgb=WHITE if item.get("type") in ("cover","closing") else NAVY
        body=slide.shapes.add_textbox(Inches(1),Inches(1.7),Inches(11.2),Inches(4.9)); tf=body.text_frame; tf.word_wrap=True; tf.clear()
        for i,line in enumerate((item.get("body") or "").split("\n")):
            par=tf.paragraphs[0] if i==0 else tf.add_paragraph(); par.text=line; par.font.size=Pt(16); par.font.color.rgb=RGBColor(205,220,228) if item.get("type") in ("cover","closing") else MUTED; par.space_after=Pt(8)
        foot=slide.shapes.add_textbox(Inches(1),Inches(7.05),Inches(11),Inches(.2)); fp=foot.text_frame.paragraphs[0]; fp.text="PMC Intelligence"; fp.font.size=Pt(8); fp.font.color.rgb=accent
    out=BytesIO(); prs.save(out); out.seek(0); return out
