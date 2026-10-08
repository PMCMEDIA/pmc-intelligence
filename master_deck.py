"""PMC presentation reference library and structured slide planning.

The uploaded master remains unchanged. This module provides a controlled,
auditable slide selection plan and protects fixed agency copy.
"""
from pathlib import Path
from datetime import datetime, timezone
from pptx import Presentation
import os, json

MASTER_DIR=Path(os.getenv("PMC_MASTER_DIR","/var/data/pmc_templates" if os.path.isdir("/var/data") else "pmc_templates"))
MASTER_PATH=MASTER_DIR/"PMC_Master_Strategy.pptx"
MAX_MASTER_BYTES=25*1024*1024

# Reference slide numbers from the approved DeQuattro + Viva presentation.
# Multiple-client pages are examples, not reusable unchanged content.
SLIDE_LIBRARY=[
 {"key":"cover","slide":1,"title":"Cover","kind":"dynamic","always":True},
 {"key":"agency_intro","slide":2,"title":"PMC Positioning","kind":"fixed","always":True},
 {"key":"agency_locations","slide":3,"title":"PMC Locations","kind":"fixed","always":True},
 {"key":"agency_values","slide":4,"title":"PMC Approach","kind":"fixed","always":True},
 {"key":"executive","slide":5,"title":"Executive Summary","kind":"dynamic","always":True},
 {"key":"production_divider","slide":6,"title":"Cinematic Production","kind":"divider","department":"Production"},
 {"key":"production_strategy","slide":7,"title":"Cinematic Strategy","kind":"dynamic","department":"Production"},
 {"key":"digital_divider","slide":9,"title":"Digital Disruption","kind":"divider","department":"Digital Media"},
 {"key":"digital_funnel","slide":11,"title":"Integrated Marketing Funnel","kind":"dynamic","department":"Digital Media"},
 {"key":"geography","slide":12,"title":"Geographic Focus","kind":"dynamic","department":"Digital Media"},
 {"key":"audience","slide":13,"title":"Audience Focus","kind":"dynamic","department":"Digital Media"},
 {"key":"competition","slide":14,"title":"Competitive Landscape","kind":"dynamic","department":"Digital Media"},
 {"key":"search","slide":15,"title":"Google Search","kind":"dynamic","department":"Digital Media"},
 {"key":"pmax","slide":16,"title":"Performance Max","kind":"dynamic","department":"Digital Media"},
 {"key":"paid_social","slide":17,"title":"Paid Social","kind":"dynamic","department":"Digital Media"},
 {"key":"ctv","slide":18,"title":"Performance Media","kind":"dynamic","department":"Digital Media"},
 {"key":"seo","slide":22,"title":"SEO Strategy","kind":"dynamic","department":"SEO / AEO"},
 {"key":"seo_ai","slide":23,"title":"AI-assisted SEO","kind":"dynamic","department":"SEO / AEO"},
 {"key":"social_divider","slide":24,"title":"Social Media Impact","kind":"divider","department":"Organic Social"},
 {"key":"social_strategy","slide":25,"title":"Social Media Management","kind":"dynamic","department":"Organic Social"},
 {"key":"social_deliverables","slide":26,"title":"Social Cadence & Platforms","kind":"dynamic","department":"Organic Social"},
 {"key":"creative_divider","slide":28,"title":"Design Revolution","kind":"divider","department":"Creative / Website"},
 {"key":"creative_strategy","slide":29,"title":"Creative Strategy","kind":"dynamic","department":"Creative / Website"},
 {"key":"website_divider","slide":31,"title":"Website Experience","kind":"divider","department":"Creative / Website"},
 {"key":"website_strategy","slide":33,"title":"Website Recommendations","kind":"dynamic","department":"Creative / Website"},
 {"key":"investment","slide":82,"title":"Investment Framework","kind":"dynamic","always":True},
 {"key":"closing","slide":86,"title":"Closing","kind":"fixed","always":True},
]
def master_status():
    if not MASTER_PATH.is_file(): return {"installed":False,"slide_count":0,"message":"Upload the approved PMC PowerPoint master."}
    try:
        prs=Presentation(str(MASTER_PATH))
        return {"installed":True,"slide_count":len(prs.slides),"updated_at":datetime.fromtimestamp(MASTER_PATH.stat().st_mtime,timezone.utc).isoformat()}
    except Exception: return {"installed":False,"slide_count":0,"message":"Master file is invalid."}

def install_master(upload):
    MASTER_DIR.mkdir(parents=True,exist_ok=True)
    data=upload.read(MAX_MASTER_BYTES+1)
    if len(data)>MAX_MASTER_BYTES: raise ValueError("Master PowerPoint exceeds 25 MB.")
    if not data.startswith(b"PK"): raise ValueError("Expected a .pptx file.")
    from io import BytesIO
    deck=Presentation(BytesIO(data))
    if len(deck.slides)<10: raise ValueError("The master must have at least 10 slides.")
    temp=MASTER_DIR/"master_pending.pptx"
    temp.write_bytes(data);temp.replace(MASTER_PATH)
    return master_status()

def proposal_plan(project):
    departments=project.get("strategy",{}).get("departments",{}) or {}
    present={k for k,v in departments.items() if v}
    result=[]
    for spec in SLIDE_LIBRARY:
        include=bool(spec.get("always") or spec.get("department") in present)
        # Specific digital channels require a corresponding recommendation.
        if include and spec["key"] in ("search","pmax","paid_social","ctv"):
            tactics=" ".join(departments.get("Digital Media",[])).lower()
            hints={"search":["search"],"pmax":["performance max","pmax"],"paid_social":["paid social","meta"],"ctv":["ctv","performance media","streaming"]}
            include=any(h in tactics for h in hints[spec["key"]])
        result.append({**spec,"include":include,"review_required":spec["kind"]=="dynamic"})
    return {"slides":result,"source":"PMC reference presentation","master":master_status(),
            "notes":["Preserve agency-approved copy on fixed slides.","Dynamic slides require client-specific copy and asset review.","Do not reuse DeQuattro/Viva-specific copy, logos, figures or claims for other clients.","Only verified benchmarks and documented financial assumptions may appear in ROI slides."]}
